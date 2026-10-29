from __future__ import annotations

import json
from contextlib import closing
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .database import Database, as_float, as_int
from .diagnostic_rules import validate_rules


MONITOR_COLUMNS = [
    "id",
    "name",
    "monitor_type",
    "enabled",
    "url",
    "method",
    "interval_seconds",
    "timeout_seconds",
    "expected_status",
    "expected_text",
    "expected_json_path",
    "expected_json_value",
    "headers_json",
    "body",
    "api_token",
    "api_token_header",
    "api_token_prefix",
    "auth_login_url",
    "auth_username",
    "auth_password",
    "auth_username_field",
    "auth_password_field",
    "auth_extra_fields",
    "auth_verify_url",
    "auth_verify_text",
    "retry_count",
    "retry_delay_seconds",
    "cooldown_minutes",
    "language",
    "web_server",
    "network_profile",
    "diagnostic_rules_json",
]

TYPE_ALIASES = {
    "web": "simple_web",
    "simple": "simple_web",
    "simple_web": "simple_web",
    "auth": "auth_web",
    "auth_web": "auth_web",
    "api": "api",
}

METHODS = {"GET", "POST", "PUT", "PATCH", "DELETE"}


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def summary(self) -> str:
        lines = [
            f"Created: {self.created}",
            f"Updated: {self.updated}",
            f"Skipped: {self.skipped}",
        ]
        if self.errors:
            lines.append("Errors:")
            lines.extend(f"- {error}" for error in self.errors[:12])
            if len(self.errors) > 12:
                lines.append(f"- ...and {len(self.errors) - 12} more")
        return "\n".join(lines)


def create_monitor_template(path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Monitors"
    ws.append(MONITOR_COLUMNS)

    examples = [
        {
            "name": "Homepage",
            "language": "PHP",
            "web_server": "Apache",
            "monitor_type": "simple_web",
            "enabled": 0,
            "url": "https://example.com",
            "method": "GET",
            "interval_seconds": 60,
            "timeout_seconds": 10,
            "expected_status": 200,
            "expected_text": "Example Domain",
            "headers_json": "{}",
            "retry_count": 3,
            "retry_delay_seconds": 2,
            "cooldown_minutes": 30,
        },
        {
            "name": "Health API",
            "language": "Python",
            "web_server": "Nginx",
            "monitor_type": "api",
            "enabled": 0,
            "url": "https://api.example.com/health",
            "method": "GET",
            "interval_seconds": 60,
            "timeout_seconds": 10,
            "expected_status": 200,
            "expected_json_path": "status",
            "expected_json_value": "ok",
            "diagnostic_rules_json": json.dumps([{"id": "api-failure", "enabled": True, "priority": 100, "phase": "response", "conditions": [{"source": "json", "path": "success", "operator": "equals", "value": False}], "result": {"error_type": "assertion_failed", "reason_code": "business_rule_match"}}]),
            "headers_json": '{"Accept":"application/json"}',
            "api_token_header": "Authorization",
            "api_token_prefix": "Bearer",
            "retry_count": 3,
            "retry_delay_seconds": 2,
            "cooldown_minutes": 30,
        },
        {
            "name": "Member Portal",
            "language": "ASP.NET",
            "web_server": "IIS",
            "monitor_type": "auth_web",
            "enabled": 0,
            "url": "https://example.com/dashboard",
            "method": "GET",
            "interval_seconds": 120,
            "timeout_seconds": 15,
            "expected_status": 200,
            "headers_json": "{}",
            "auth_login_url": "https://example.com/login",
            "auth_username": "demo@example.com",
            "auth_password": "change-me",
            "auth_username_field": "username",
            "auth_password_field": "password",
            "auth_extra_fields": "{}",
            "auth_verify_url": "https://example.com/dashboard",
            "auth_verify_text": "Dashboard",
            "retry_count": 3,
            "retry_delay_seconds": 2,
            "cooldown_minutes": 30,
        },
    ]
    for example in examples:
        ws.append([example.get(column, "") for column in MONITOR_COLUMNS])

    _style_template_sheet(ws)
    _add_validations(ws)
    _add_instructions(wb)
    wb.save(output)
    return output


def import_monitors_from_excel(db: Database, path: str | Path) -> ImportResult:
    with closing(load_workbook(path, read_only=True, data_only=True)) as workbook:
        return _import_workbook(db, workbook)


def _import_workbook(db: Database, wb: Any) -> ImportResult:
    result = ImportResult()
    ws = wb["Monitors"] if "Monitors" in wb.sheetnames else wb.active

    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        result.errors.append("Workbook is empty")
        return result

    headers = [_normalize_header(value) for value in header_row]
    header_map = {header: index for index, header in enumerate(headers) if header}
    missing = [column for column in ("name", "monitor_type", "url") if column not in header_map]
    if missing:
        result.errors.append(f"Missing required columns: {', '.join(missing)}")
        return result

    for row_number, row in enumerate(rows, start=2):
        values = _row_to_dict(headers, row)
        if not any(str(value or "").strip() for value in values.values()):
            result.skipped += 1
            continue
        try:
            payload = _normalize_monitor(values)
        except ValueError as exc:
            result.errors.append(f"Row {row_number}: {exc}")
            result.skipped += 1
            continue

        try:
            monitor_id = as_int(values.get("id"), 0)
            existing = db.get_monitor(monitor_id) if monitor_id else None
            existing = existing or db.find_monitor_by_name_url(payload["name"], payload["url"])
            if existing:
                payload["id"] = existing["id"]
                db.save_monitor(payload)
                result.updated += 1
            else:
                db.save_monitor(payload)
                result.created += 1
        except ValueError as exc:
            result.errors.append(f"Row {row_number}: {exc}")
            result.skipped += 1

    return result


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _row_to_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    data: dict[str, Any] = {}
    for index, header in enumerate(headers):
        if header:
            data[header] = row[index] if index < len(row) else None
    return data


def _normalize_monitor(values: dict[str, Any]) -> dict[str, Any]:
    name = str(values.get("name") or "").strip()
    url = str(values.get("url") or "").strip()
    if not name:
        raise ValueError("name is required")
    if not url:
        raise ValueError("url is required")

    monitor_type = TYPE_ALIASES.get(str(values.get("monitor_type") or "").strip().lower())
    if not monitor_type:
        raise ValueError("monitor_type must be simple_web, auth_web, or api")

    method = str(values.get("method") or "GET").strip().upper()
    if method not in METHODS:
        raise ValueError("method must be GET, POST, PUT, PATCH, or DELETE")

    headers_json = _json_object_text(values.get("headers_json"), "headers_json", default="{}")
    auth_extra_fields = _json_object_text(values.get("auth_extra_fields"), "auth_extra_fields", default="{}")

    payload = {
        "name": name,
        "monitor_type": monitor_type,
        "enabled": _bool_value(values.get("enabled"), default=True),
        "url": url,
        "method": method,
        "interval_seconds": max(10, as_int(values.get("interval_seconds"), 60)),
        "timeout_seconds": max(1.0, as_float(values.get("timeout_seconds"), 10.0)),
        "expected_status": as_int(values.get("expected_status"), 200),
        "expected_text": str(values.get("expected_text") or ""),
        "expected_json_path": str(values.get("expected_json_path") or ""),
        "expected_json_value": str(values.get("expected_json_value") or ""),
        "headers_json": headers_json,
        "body": str(values.get("body") or ""),
        "api_token": str(values.get("api_token") or ""),
        "api_token_header": str(values.get("api_token_header") or "Authorization"),
        "api_token_prefix": str(values.get("api_token_prefix") or "Bearer"),
        "auth_login_url": str(values.get("auth_login_url") or ""),
        "auth_username": str(values.get("auth_username") or ""),
        "auth_password": str(values.get("auth_password") or ""),
        "auth_username_field": str(values.get("auth_username_field") or "username"),
        "auth_password_field": str(values.get("auth_password_field") or "password"),
        "auth_extra_fields": auth_extra_fields,
        "auth_verify_url": str(values.get("auth_verify_url") or ""),
        "auth_verify_text": str(values.get("auth_verify_text") or ""),
        "retry_count": max(1, as_int(values.get("retry_count"), 3)),
        "retry_delay_seconds": max(0.0, as_float(values.get("retry_delay_seconds"), 2.0)),
        "cooldown_minutes": max(0, as_int(values.get("cooldown_minutes"), 30)),
        **{key: str(values.get(key) or "").strip() for key in ("language", "web_server") if key in values},
    }
    profile = str(values.get("network_profile") or "").strip()
    if profile:
        payload["network_profile"] = "" if profile == "__CLEAR__" else profile
    raw_rules = str(values.get("diagnostic_rules_json") or "").strip()
    if raw_rules:
        try:
            payload["diagnostic_rules_json"] = json.dumps(validate_rules(raw_rules), ensure_ascii=False)
        except ValueError as exc:
            raise ValueError(f"diagnostic_rules_json: {exc}") from exc
    return payload


def _bool_value(value: Any, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "on", "enabled"}:
        return True
    if text in {"0", "false", "no", "n", "off", "disabled"}:
        return False
    return default


def _json_object_text(value: Any, label: str, default: str = "{}") -> str:
    text = str(value or "").strip()
    if not text:
        return default
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{label} is invalid JSON: {exc}") from exc
    if not isinstance(parsed, dict):
        raise ValueError(f"{label} must be a JSON object")
    return json.dumps(parsed, ensure_ascii=False)


def _style_template_sheet(ws: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(MONITOR_COLUMNS)).coordinate}"
    widths = {
        "A": 8,
        "B": 24,
        "C": 16,
        "D": 10,
        "E": 36,
        "G": 16,
        "H": 16,
        "M": 34,
        "N": 34,
        "O": 28,
        "R": 34,
        "X": 34,
        "AE": 24,
        "AF": 58,
    }
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        ws.column_dimensions[letter].width = widths.get(letter, 18)
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 84 if row[0].row == 3 else 42
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_validations(ws: Any) -> None:
    type_validation = DataValidation(type="list", formula1='"simple_web,auth_web,api"', allow_blank=False)
    method_validation = DataValidation(type="list", formula1='"GET,POST,PUT,PATCH,DELETE"', allow_blank=True)
    enabled_validation = DataValidation(type="list", formula1='"1,0,TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(type_validation)
    ws.add_data_validation(method_validation)
    ws.add_data_validation(enabled_validation)
    type_validation.add("C2:C500")
    enabled_validation.add("D2:D500")
    method_validation.add("F2:F500")


def _add_instructions(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions")
    rows = [
        ("Column", "Required", "Notes"),
        ("id", "No", "Existing monitor id. Leave blank to create. If id is blank, import updates by exact name + url match."),
        ("name", "Yes", "Monitor display name."),
        ("monitor_type", "Yes", "simple_web, auth_web, or api."),
        ("enabled", "No", "1/TRUE/yes to enable, 0/FALSE/no to disable."),
        ("url", "Yes", "Target website, protected page, or API endpoint."),
        ("language", "No", "Application language/framework, for example PHP, Python, Java, ASP.NET."),
        ("web_server", "No", "Web server, for example Apache, Nginx, IIS, Tomcat."),
        ("network_profile", "No", "Existing network profile; blank preserves it; __CLEAR__ removes it."),
        ("diagnostic_rules_json", "No", "JSON array of rules; blank preserves existing rules; [] clears them. See the Health API example and checker-spec.md §7. Configure profiles first."),
        ("headers_json", "No", "JSON object, for example {\"Accept\":\"application/json\"}."),
        ("body", "No", "Request body for API methods such as POST/PUT/PATCH."),
        ("api_token", "No", "Stored encrypted after import."),
        ("auth_password", "No", "Stored encrypted after import."),
        ("auth_extra_fields", "No", "JSON object with additional login form fields. Hidden inputs are collected from the login page automatically."),
    ]
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 100
    for row in ws.iter_rows():
        ws.row_dimensions[row[0].row].height = (56 if row[0].value == "diagnostic_rules_json" else 34) if row[0].row > 1 else 24
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
