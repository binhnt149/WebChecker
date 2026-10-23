from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import Database


def format_duration(seconds: float) -> str:
    seconds = max(0, round(seconds))
    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def build_report(db: Database, start: datetime, end: datetime, now: datetime | None = None) -> dict[str, Any]:
    if start > end:
        raise ValueError("From must be before To")
    effective_end = min(end, now or datetime.now())
    monitors = db.list_monitors(include_disabled=True)
    rows: dict[tuple[str, Any], dict[str, Any]] = {}
    for monitor in monitors:
        rows[("id", monitor["id"])] = {
            "monitor_id": monitor["id"], "monitor_name": monitor["name"],
            "language": monitor.get("language", ""), "web_server": monitor.get("web_server", ""),
            "down_count": 0, "incident_count": 0, "downtime_seconds": 0.0,
            "average_down_seconds": 0.0, "open_incidents": 0, "last_status": "unknown",
        }
    active: dict[tuple[str, Any], tuple[datetime, bool]] = {}
    incidents: list[dict[str, Any]] = []

    def close_incident(key: tuple[str, Any], until: datetime, is_open: bool) -> None:
        began, carried = active.pop(key)
        counted_from = max(began, start)
        counted_until = min(until, effective_end)
        if counted_until < counted_from or (carried and counted_until == counted_from):
            return
        seconds = (counted_until - counted_from).total_seconds()
        row = rows[key]
        row["incident_count"] += 1
        row["downtime_seconds"] += seconds
        row["open_incidents"] += int(is_open)
        incidents.append({
            "monitor_id": row["monitor_id"], "monitor_name": row["monitor_name"],
            "counted_from": counted_from.isoformat(timespec="seconds"),
            "counted_until": counted_until.isoformat(timespec="seconds"),
            "recovered_at": None if is_open else until.isoformat(timespec="seconds"),
            "duration_seconds": seconds, "carried_in": carried, "open": is_open,
        })

    logs = db.report_logs(start.isoformat(timespec="seconds"), effective_end.isoformat(timespec="seconds")) if effective_end >= start else []
    for log in logs:
        key = ("id", log["monitor_id"]) if log["monitor_id"] is not None else ("name", log["monitor_name"])
        if key not in rows:
            rows[key] = {
                "monitor_id": log["monitor_id"], "monitor_name": log["monitor_name"], "language": "", "web_server": "",
                "down_count": 0, "incident_count": 0, "downtime_seconds": 0.0,
                "average_down_seconds": 0.0, "open_incidents": 0, "last_status": "unknown",
            }
        row = rows[key]
        checked_at = datetime.fromisoformat(log["checked_at"])
        state = log["status"]
        if state == "down":
            if checked_at >= start:
                row["down_count"] += 1
            if key not in active:
                active[key] = (max(checked_at, start), checked_at < start)
        elif state == "up" and key in active:
            close_incident(key, checked_at, False)
        if state in {"up", "down"}:
            row["last_status"] = state
    for key in list(active):
        close_incident(key, effective_end, True)
    for row in rows.values():
        if row["incident_count"]:
            row["average_down_seconds"] = row["downtime_seconds"] / row["incident_count"]
    values = sorted(rows.values(), key=lambda row: row["monitor_name"].casefold())
    current_ids = {monitor["id"] for monitor in monitors}
    current = [row for row in values if row["monitor_id"] in current_ids]
    return {
        "start_at": start.isoformat(timespec="seconds"), "end_at": end.isoformat(timespec="seconds"),
        "as_of": effective_end.isoformat(timespec="seconds"), "rows": values, "incidents": incidents,
        "total": len(current), "up": sum(row["last_status"] == "up" for row in current),
        "down": sum(row["last_status"] == "down" for row in current),
        "unknown": sum(row["last_status"] == "unknown" for row in current),
        "down_count": sum(row["down_count"] for row in values),
        "downtime_seconds": sum(row["downtime_seconds"] for row in values),
    }


def export_report(report: dict[str, Any], path: str | Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Statistics"
    summary.append(["SiteSentry downtime report"])
    summary.append(["From", datetime.fromisoformat(report["start_at"]), "", "To", datetime.fromisoformat(report["end_at"])])
    summary.append(["Observed through", datetime.fromisoformat(report["as_of"])])
    summary.append(["DOWN count = DOWN log rows. Incident = consecutive DOWNs until UP. Durations are clipped to the selected period."])
    summary.merge_cells("A1:I1")
    summary.merge_cells("A4:I4")
    summary.merge_cells("B2:C2")
    summary.merge_cells("E2:F2")
    summary.merge_cells("B3:C3")
    summary["A1"].font = Font(size=18, bold=True, color="1F2937")
    summary.row_dimensions[1].height = 30
    summary.row_dimensions[2].height = 34
    summary.row_dimensions[3].height = 34
    summary.row_dimensions[4].height = 32
    summary.append(["Monitor name", "DOWN count", "Incidents", "Total DOWN", "Average DOWN / incident", "Open incidents", "Language", "Web server", "Last status"])
    for row in report["rows"]:
        summary.append([
            row["monitor_name"], row["down_count"], row["incident_count"], row["downtime_seconds"] / 86400,
            row["average_down_seconds"] / 86400, row["open_incidents"], row["language"], row["web_server"], row["last_status"],
        ])
    for cell in (summary["B2"], summary["E2"], summary["B3"]):
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for row in summary.iter_rows(min_row=6):
        row[3].number_format = row[4].number_format = "[h]:mm:ss"
    details = workbook.create_sheet("Incidents")
    details.append(["Monitor name", "Counted from", "Counted until", "Recovered at", "Duration", "Carried in", "Still DOWN"])
    for incident in report["incidents"]:
        details.append([
            incident["monitor_name"], datetime.fromisoformat(incident["counted_from"]),
            datetime.fromisoformat(incident["counted_until"]),
            datetime.fromisoformat(incident["recovered_at"]) if incident["recovered_at"] else None,
            incident["duration_seconds"] / 86400, incident["carried_in"], incident["open"],
        ])
    for row in details.iter_rows(min_row=2):
        for cell in row[1:4]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        row[4].number_format = "[h]:mm:ss"
    for sheet, header_row, widths in [
        (summary, 5, [32, 16, 14, 22, 26, 18, 20, 24, 16]),
        (details, 1, [32, 24, 24, 24, 20, 14, 14]),
    ]:
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(max(header_row, sheet.max_row), len(widths)).coordinate}"
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.data_type == "f":
                    cell.data_type = "s"
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="1F2937")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.row_dimensions[header_row].height = 32
    for address in ("B2", "E2", "B3"):
        summary[address].alignment = Alignment(horizontal="left", vertical="top")
    workbook.save(path)
