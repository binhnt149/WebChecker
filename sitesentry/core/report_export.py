from __future__ import annotations

from datetime import datetime
from math import ceil
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .diagnostics import error_label


INK = "18252C"
TEAL = "087D72"
PALE = "F1F6F8"
LINE = "D6E0E5"
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"


def _append(sheet: Any, values: list[Any]) -> None:
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _style(sheet: Any, header: int, widths: list[int], last_data: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = TEAL if sheet.title == "Statistics" else "52616B"
    sheet.freeze_panes = f"B{header + 1}"
    end = get_column_letter(len(widths))
    sheet.auto_filter.ref = f"A{header}:{end}{max(header, last_data)}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet:
        height = 22
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            numeric = isinstance(cell.value, (int, float, datetime))
            cell.alignment = Alignment(horizontal="right" if numeric else "left", vertical="center", wrap_text=True, indent=1 if cell.row > header else 0)
            if cell.row > header and cell.row <= last_data and cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
            if cell.row > header and cell.value is not None:
                height = max(height, 15 * ceil(len(str(cell.value)) / max(8, widths[cell.column - 1] - 2)))
        if row[0].row > header:
            sheet.row_dimensions[row[0].row].height = min(400, height)
    for cell in sheet[header]:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
    sheet.row_dimensions[header].height = 32
    sheet.print_title_rows = f"{header}:{header}"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f"A1:{end}{sheet.max_row}"
    sheet.oddFooter.left.text = "SiteSentry Desktop"
    sheet.oddFooter.right.text = "Page &P of &N"


def export_report(report: dict[str, Any], path: str | Path) -> None:
    workbook = Workbook()
    workbook.properties.creator = "SiteSentry Desktop"
    workbook.properties.title = "Monitor availability and downtime"
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    summary = workbook.active
    summary.title = "Statistics"
    _append(summary, ["SiteSentry availability report"])
    _append(summary, ["From", datetime.fromisoformat(report["start_at"]), "", "To", datetime.fromisoformat(report["end_at"])])
    _append(summary, ["Observed through", datetime.fromisoformat(report["as_of"]), "", "DOWN checks", report["down_count"], "", "Incidents", len(report["incidents"])])
    _append(summary, ["DOWN checks count log observations. Average downtime is per incident. Open incidents are measured through the observation time."])
    summary.merge_cells("A1:M1")
    summary.merge_cells("B2:C2")
    summary.merge_cells("E2:F2")
    summary.merge_cells("B3:C3")
    summary.merge_cells("A4:M4")
    _append(summary, ["Monitor name", "DOWN checks", "Incidents", "Total DOWN", "Average DOWN", "Open incidents", "Language", "Web server", "Last status", "Current error category", "Classification", "Reason code", "Phase"])
    for row in report["rows"]:
        _append(summary, [row["monitor_name"], row["down_count"], row["incident_count"], row["downtime_seconds"] / 86400, row["average_down_seconds"] / 86400, row["open_incidents"], row["language"], row["web_server"], row["last_status"].upper(), error_label(row.get("last_error_type", "")), row.get("last_classification_state", ""), row.get("last_reason_code", ""), row.get("last_phase", "")])
    last_data = summary.max_row
    _style(summary, 5, [34, 15, 13, 20, 20, 16, 18, 20, 14, 36, 20, 25, 20], last_data)
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color=INK)
    summary["A1"].border = Border(bottom=Side(style="thin", color=TEAL))
    summary["A4"].font = Font(name="Arial", size=10, italic=True, color="52616B")
    for row, height in [(1, 32), (2, 28), (3, 28), (4, 32)]:
        summary.row_dimensions[row].height = height
    for address in ("B2", "E2", "B3"):
        summary[address].number_format = DATE_FORMAT
        summary[address].alignment = Alignment(horizontal="left", vertical="center")
    for row in summary.iter_rows(min_row=6, max_row=last_data):
        row[3].number_format = row[4].number_format = "[h]:mm:ss"
        row[5].alignment = Alignment(horizontal="center", vertical="center")
    if last_data >= 6:
        summary.conditional_formatting.add(f"F6:F{last_data}", CellIsRule(operator="greaterThan", formula=[0], fill=PatternFill("solid", fgColor="FDE7EA"), font=Font(color="B42332", bold=True)))
        total_row = last_data + 2
        summary.cell(total_row, 1, "Filtered total")
        for column in (2, 3, 4, 6):
            letter = get_column_letter(column)
            summary.cell(total_row, column, f"=SUBTOTAL(109,{letter}6:{letter}{last_data})")
        summary.cell(total_row, 5, f"=IF(C{total_row}=0,0,D{total_row}/C{total_row})")
        for cell in summary[total_row]:
            cell.font = Font(name="Arial", size=10, bold=True, color=INK)
            cell.border = Border(top=Side(style="thin", color=LINE))
        summary.cell(total_row, 4).number_format = summary.cell(total_row, 5).number_format = "[h]:mm:ss"
    details = workbook.create_sheet("Incidents")
    _append(details, ["Monitor name", "Counted from", "Counted until", "Recovered at", "Duration", "Carried in", "Still DOWN", "Error categories", "Classification", "Reason codes"])
    for incident in report["incidents"]:
        _append(details, [incident["monitor_name"], datetime.fromisoformat(incident["counted_from"]), datetime.fromisoformat(incident["counted_until"]), datetime.fromisoformat(incident["recovered_at"]) if incident["recovered_at"] else None, incident["duration_seconds"] / 86400, "Yes" if incident["carried_in"] else "No", "Yes" if incident["open"] else "No", ", ".join(error_label(code) for code in incident.get("error_types", [])), ", ".join(incident.get("classification_states", [])), ", ".join(incident.get("reason_codes", []))])
    _style(details, 1, [34, 24, 24, 24, 20, 13, 13, 48, 22, 36], details.max_row)
    for row in details.iter_rows(min_row=2):
        for cell in row[1:4]:
            cell.number_format = DATE_FORMAT
        row[4].number_format = "[h]:mm:ss"
        for cell in row[5:7]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    errors = workbook.create_sheet("Error breakdown")
    _append(errors, ["Monitor name", "Error code", "Error category", "DOWN checks", "Classification", "Reason code", "Phase"])
    for row in report["rows"]:
        counts = row.get("diagnostic_counts")
        if counts is None:
            counts = [{"error_type": code, "count": count} for code, count in sorted(row.get("error_counts", {}).items())]
        for item in counts:
            code = item["error_type"]
            _append(errors, [row["monitor_name"], code, error_label(code), item["count"], item.get("classification_state", ""), item.get("reason_code", ""), item.get("phase", "")])
    _style(errors, 1, [34, 26, 48, 18, 20, 25, 20], errors.max_row)
    summary.print_area = f"A1:M{summary.max_row}"
    workbook.save(path)


def create_report_template(path: str | Path, start: datetime | None = None, end: datetime | None = None) -> None:
    now = datetime.now().replace(microsecond=0)
    start = start or now.replace(hour=0, minute=0, second=0)
    end = end or now
    export_report({"start_at": start.isoformat(), "end_at": end.isoformat(), "as_of": min(end, now).isoformat(), "rows": [], "incidents": [], "down_count": 0}, path)
