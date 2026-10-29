from __future__ import annotations

import csv
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from typing import Any

import customtkinter as ctk
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .background import run_background
from .datetime_picker import DateTimePicker
from .focus import bind_focus_on_click
from .theme import configure_table
from sitesentry.core.diagnostics import ERROR_TYPES, error_label


EXPORT_FIELDS = [
    "id", "monitor_id", "monitor_name", "status", "ok", "response_ms",
    "status_code", "error_type", "error", "attempts", "checked_at", "detail_json",
    "classification_state", "reason_code", "phase", "evidence_source", "evidence_json",
]


class LogsView(ctk.CTkFrame):
    def __init__(self, master: Any, app: Any) -> None:
        super().__init__(master, fg_color="transparent")
        self.app = app
        self._busy = False
        self._refresh_again = False
        self._logs: dict[str, dict[str, Any]] = {}
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        toolbar = ctk.CTkFrame(self, fg_color="transparent")
        toolbar.grid(row=0, column=0, padx=20, pady=(18, 8), sticky="ew")
        toolbar.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(toolbar, text="Logs", font=ctk.CTkFont(size=24, weight="bold")).grid(row=0, column=0, sticky="w")
        actions = ctk.CTkFrame(toolbar, fg_color="transparent")
        actions.grid(row=0, column=1, sticky="e")
        ctk.CTkButton(actions, text="Export Excel", command=lambda: self._export("xlsx"), width=110).pack(side="left", padx=4)
        ctk.CTkButton(actions, text="Export CSV", command=lambda: self._export("csv"), width=100, fg_color="#52525b").pack(side="left", padx=4)
        ctk.CTkButton(actions, text="Clear", command=self._clear_logs, width=64, fg_color="#991b1b").pack(side="left", padx=(4, 0))
        filters = ctk.CTkFrame(toolbar, fg_color="transparent")
        filters.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(14, 0))
        filters.grid_columnconfigure((0, 1), weight=1)
        self.start_picker = DateTimePicker(filters, "From")
        self.start_picker.grid(row=0, column=0, padx=(0, 12), sticky="ew")
        self.end_picker = DateTimePicker(filters, "To", end_of_day=True)
        self.end_picker.grid(row=0, column=1, sticky="ew")
        status_row = ctk.CTkFrame(toolbar, fg_color="transparent")
        status_row.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(10, 0))
        self.status_filter = tk.StringVar(value="all")
        ctk.CTkLabel(status_row, text="Status").pack(side="left", padx=(0, 8))
        ctk.CTkOptionMenu(status_row, values=["all", "up", "down"], variable=self.status_filter, command=lambda _: self.refresh(), width=100).pack(side="left")
        self.error_filter = tk.StringVar(value="All errors")
        ctk.CTkOptionMenu(status_row, values=["All errors", *ERROR_TYPES.values()], variable=self.error_filter, command=lambda _: self.refresh(), width=280).pack(side="left", padx=12)
        ctk.CTkButton(status_row, text="Apply", command=self.refresh, width=80).pack(side="right")
        detail_filters = ctk.CTkFrame(toolbar, fg_color="transparent")
        detail_filters.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        self.monitor_filter = tk.StringVar()
        self.classification_filter = tk.StringVar(value="all")
        self.phase_filter = tk.StringVar(value="all")
        bind_focus_on_click(ctk.CTkEntry(detail_filters, textvariable=self.monitor_filter, placeholder_text="Monitor name", width=180)).pack(side="left", padx=(0, 8))
        ctk.CTkOptionMenu(detail_filters, values=["all", "confirmed", "undetermined", "legacy", "not_applicable"], variable=self.classification_filter, width=155, command=lambda _: self.refresh()).pack(side="left", padx=(0, 8))
        ctk.CTkOptionMenu(detail_filters, values=["all", "configuration", "resolve", "connect", "tls", "request", "login", "login_page", "login_submit", "verify", "response", "assertion", "network_probe"], variable=self.phase_filter, width=140, command=lambda _: self.refresh()).pack(side="left")
        table_frame = ctk.CTkFrame(self, fg_color="transparent")
        table_frame.grid(row=1, column=0, padx=20, pady=(8, 12), sticky="nsew")
        table_frame.grid_columnconfigure(0, weight=1)
        table_frame.grid_rowconfigure(0, weight=1)
        columns = ("monitor_name", "status", "checked_at", "error_type", "classification_state", "reason_code", "phase", "response_ms", "status_code", "attempts", "error")
        self.table = ttk.Treeview(table_frame, columns=columns, show="headings", style="SiteSentry.Logs.Treeview")
        for column, title, width in zip(columns, ["Monitor", "Status", "Checked at", "Error category", "Classification", "Reason", "Phase", "Time (ms)", "HTTP", "Attempts", "Error"], [220, 80, 175, 260, 130, 160, 120, 100, 70, 80, 420]):
            self.table.heading(column, text=title)
            self.table.column(column, width=width, minwidth=60, anchor="w")
        dark = ctk.get_appearance_mode() == "Dark"
        ttk.Style(self).configure("SiteSentry.Logs.Treeview", rowheight=30, font=("Arial", 12), background="#242424" if dark else "#ffffff", fieldbackground="#242424" if dark else "#ffffff", foreground="#eeeeee" if dark else "#222222")
        self.table.tag_configure("down", foreground="#dc2626")
        self.table.grid(row=0, column=0, sticky="nsew")
        vertical = ttk.Scrollbar(table_frame, orient="vertical", command=self.table.yview)
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal = ttk.Scrollbar(table_frame, orient="horizontal", command=self.table.xview)
        horizontal.grid(row=1, column=0, sticky="ew")
        self.table.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        configure_table(self.table)
        self.table.bind("<Double-1>", lambda _: self._show_details())
        self.summary_label = ctk.CTkLabel(self, text="", anchor="w")
        self.summary_label.grid(row=2, column=0, padx=20, pady=(0, 12), sticky="ew")
        self.refresh()

    def _date_range(self) -> tuple[str | None, str | None] | None:
        start, end = self.start_picker.get(), self.end_picker.get()
        if start and end and start > end:
            messagebox.showerror("Log date filter", "From must be before To.", parent=self)
            return None
        return start.isoformat(timespec="seconds") if start else None, end.isoformat(timespec="seconds") if end else None

    def refresh(self) -> None:
        if self._busy:
            self._refresh_again = True
            return
        period = self._date_range()
        if period is None:
            return
        self._busy = True
        status = self.status_filter.get()
        error_type = next((code for code, label in ERROR_TYPES.items() if label == self.error_filter.get()), None)
        filters = {"classification_state": self.classification_filter.get(), "phase": self.phase_filter.get(), "monitor_name": self.monitor_filter.get().strip()}

        def done(logs: list[dict[str, Any]]) -> None:
            self._busy = False
            children = self.table.get_children()
            if children:
                self.table.delete(*children)
            self._logs = {str(item["id"]): item for item in logs}
            for index, item in enumerate(logs):
                values = tuple(error_label(item.get(field, "")) if field == "error_type" else ("-" if item.get(field) is None else item.get(field, "")) for field in self.table["columns"])
                self.table.insert("", "end", iid=str(item["id"]), values=values, tags=(item["status"], "stripe") if index % 2 else (item["status"],))
            self.summary_label.configure(text=f"Showing {len(logs)} latest log row(s)")
            if self._refresh_again:
                self._refresh_again = False
                self.refresh()

        def failed(message: str) -> None:
            self._busy = False
            self.summary_label.configure(text=message)

        run_background(self.app, self, lambda: self.app.db.list_logs(limit=500, status_filter=status, start_at=period[0], end_at=period[1], error_type=error_type, **filters), done, failed)

    def _show_details(self) -> None:
        selected = self.table.selection()
        if not selected or selected[0] not in self._logs:
            return
        item = self._logs[selected[0]]
        dialog = ctk.CTkToplevel(self)
        dialog.title("Diagnostic evidence")
        dialog.geometry("680x440")
        text = ctk.CTkTextbox(dialog, wrap="word")
        text.pack(fill="both", expand=True, padx=16, pady=16)
        text.insert("1.0", f"{item['monitor_name']}\n{item['classification_state']} | {item['error_type']} | {item['reason_code']}\nPhase: {item['phase']} | Source: {item['evidence_source']}\n\n{item['error']}\n\n{item['evidence_json']}")
        text.configure(state="disabled")

    def _export(self, extension: str) -> None:
        period = self._date_range()
        if period is None:
            return
        path = filedialog.asksaveasfilename(parent=self, title="Export logs", initialfile=self._default_export_name(extension), defaultextension=f".{extension}", filetypes=[("Excel workbook" if extension == "xlsx" else "CSV files", f"*.{extension}")])
        if not path:
            return
        status = self.status_filter.get()
        error_type = next((code for code, label in ERROR_TYPES.items() if label == self.error_filter.get()), None)
        filters = {"classification_state": self.classification_filter.get(), "phase": self.phase_filter.get(), "monitor_name": self.monitor_filter.get().strip()}

        def work() -> int:
            logs = self.app.db.list_logs(limit=-1, status_filter=status, start_at=period[0], end_at=period[1], error_type=error_type, **filters)
            if extension == "xlsx":
                self._write_excel(logs, path)
            else:
                self._write_csv(logs, path)
            return len(logs)

        run_background(self.app, self, work, lambda count: messagebox.showinfo("Export logs", f"Exported {count} log row(s).", parent=self))

    @staticmethod
    def _write_csv(logs: list[dict[str, Any]], path: str) -> None:
        with open(path, "w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=EXPORT_FIELDS, extrasaction="ignore")
            writer.writeheader()
            for item in logs:
                row = {key: ("'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value) for key, value in item.items()}
                writer.writerow(row)

    @staticmethod
    def _write_excel(logs: list[dict[str, Any]], path: str) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Logs"
        sheet.append(EXPORT_FIELDS)
        for item in logs:
            sheet.append([item.get(field, "") for field in EXPORT_FIELDS])
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F2937")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate([8, 12, 28, 12, 8, 14, 14, 26, 50, 10, 24, 50, 20, 26, 20, 20, 60], start=1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.data_type == "f":
                    cell.data_type = "s"
        workbook.save(path)

    def _default_export_name(self, extension: str) -> str:
        values = [picker.get().strftime("%Y%m%d_%H%M%S") if picker.get() else fallback for picker, fallback in [(self.start_picker, "begin"), (self.end_picker, "now")]]
        return f"sitesentry_logs_{values[0]}_{values[1]}.{extension}"

    def _clear_logs(self) -> None:
        if messagebox.askyesno("Clear logs", "Delete all log rows? Statistics based on these logs will also be lost.", parent=self):
            self.app.db.clear_logs()
            self.refresh()
