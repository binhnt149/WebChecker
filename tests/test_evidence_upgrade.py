import tempfile
import unittest
import json
import asyncio
import smtplib
import sqlite3
from unittest.mock import patch

import httpx
from datetime import datetime, timedelta
from pathlib import Path

from sitesentry.core.checker import CheckResult
from sitesentry.core.checker import SiteChecker
from sitesentry.core.database import Database, SCHEMA_VERSION
from sitesentry.core.reports import build_report, filter_email_summary
from sitesentry.core.notifier import EmailNotifier
from sitesentry.core.security import SecurityManager
from sitesentry.core.report_export import export_report
from sitesentry.ui.logs_view import LogsView
from openpyxl import load_workbook


class UnknownIncidentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.db = Database(root / "db.sqlite", SecurityManager(root / "key"))
        self.db.set_setting("smtp_enabled", "1")
        rule = self.db.list_email_rules()[0]
        self.db.save_email_rule({**rule, "batch_window_seconds": 0})
        self.id = self.db.save_monitor({"name": "Unknown", "url": "https://example.test"})

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def record(self, kind, ok=False):
        self.db.record_check_result(self.id, CheckResult(ok, "up" if ok else "down", error_type=kind, error="synthetic", evidence_source="transport", evidence={"exception_kind": "SyntheticError"} if kind and kind != "unknown" else {}))

    def test_unknown_is_logged_without_down_or_recovery_email(self):
        self.db.save_email_rule({"name": "Recovery", "rule_type": "recovery_only"})
        self.record("unknown")
        self.record("", ok=True)
        self.assertEqual(len(self.db.list_logs()), 2)
        self.assertEqual(self.db.email_delivery_counts(), {})

    def test_unknown_becomes_confirmed_and_queues_once(self):
        self.record("unknown")
        self.assertEqual(self.db.email_delivery_counts(), {})
        self.record("timeout")
        self.record("timeout")
        self.assertEqual(self.db.email_delivery_counts(), {"pending": 1})

    def test_unknown_is_always_omitted_from_scheduled_summary(self):
        self.record("unknown")
        now = datetime.now()
        report = build_report(self.db, now - timedelta(minutes=1), now + timedelta(minutes=1), now=now)
        self.assertEqual(len(report["rows"]), 1)
        filtered = filter_email_summary(report, set())
        self.assertEqual(filtered["rows"], [])
        self.assertEqual(filtered["total"], 0)

    def test_known_down_then_unknown_still_gets_recovery(self):
        self.record("timeout")
        self.record("unknown")
        self.record("", ok=True)
        self.assertEqual(self.db.email_delivery_counts(), {"pending": 2})

    def test_unconfirmed_diagnosis_with_nonunknown_code_does_not_send(self):
        self.db.record_check_result(self.id, CheckResult(False, "down", error_type="timeout", classification_state="undetermined"))
        self.assertEqual(self.db.email_delivery_counts(), {})
        self.assertEqual(self.db.list_logs()[0]["error_type"], "unknown")

    def test_confirmed_label_without_any_evidence_is_held(self):
        self.db.record_check_result(self.id, CheckResult(False, "down", error_type="timeout", classification_state="confirmed"))
        self.assertEqual(self.db.email_delivery_counts(), {})
        self.assertEqual(self.db.list_logs()[0]["classification_state"], "undetermined")

    def test_internet_event_changes_are_recorded_once_without_creating_incident_mail(self):
        def check(state):
            result = CheckResult(True, "up", evidence={"internet_check": {"state": state, "profile": "internet", "targets": []}})
            self.db.record_check_result(self.id, result)
        check("reference_unreachable")
        check("reference_unreachable")
        check("reachable")
        self.assertEqual([event["state"] for event in self.db.list_network_events()], ["reachable", "reference_unreachable"])
        self.assertEqual(self.db.email_delivery_counts(), {})

    def test_reference_probe_alone_is_not_evidence_of_monitor_failure(self):
        result = CheckResult(False, "down", error_type="timeout", classification_state="confirmed", evidence={"internet_check": {"state": "reference_unreachable", "profile": "internet", "targets": []}})
        self.db.record_check_result(self.id, result)
        self.assertEqual(self.db.list_logs()[0]["error_type"], "unknown")
        self.assertEqual(self.db.email_delivery_counts(), {})

    def test_internet_event_ignores_older_results_and_updates_last_seen(self):
        def check(state, stamp):
            self.db.record_check_result(self.id, CheckResult(True, "up", evidence={"internet_check": {"state": state, "profile": "internet", "checked_at": stamp, "targets": []}}))
        check("reachable", "2026-01-01T10:00:00")
        check("reachable", "2026-01-01T10:01:00")
        check("reference_unreachable", "2026-01-01T09:59:00")
        events = self.db.list_network_events()
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0]["checked_at"], "2026-01-01T10:00:00")
        self.assertEqual(events[0]["last_seen_at"], "2026-01-01T10:01:00")

    def test_recovery_keeps_original_diagnosis_evidence(self):
        self.db.record_check_result(self.id, CheckResult(False, "down", error_type="timeout", reason_code="read_timeout", phase="request", evidence_source="transport", evidence={"exception_kind": "ReadTimeout"}))
        self.record("unknown")
        self.record("", ok=True)
        rows = self.db._conn.execute("SELECT payload_json FROM notification_events ORDER BY id").fetchall()
        payload = json.loads(rows[-1][0])["result"]
        self.assertEqual(payload["reason_code"], "read_timeout")
        self.assertEqual(payload["evidence"]["exception_kind"], "ReadTimeout")
        subject, body = EmailNotifier(self.db)._rule_message(self.db.pending_email_deliveries()[0])
        self.assertIn("read_timeout", body)
        self.assertIn("request", body)

    def test_unclassified_member_cannot_ride_in_confirmed_batch(self):
        rule = self.db.list_email_rules()[0]
        self.db.save_email_rule({**rule, "batch_window_seconds": 15})
        self.record("timeout")
        other = self.db.save_monitor({"name": "Other", "url": "https://other.example.test"})
        payload = {"monitor": {"id": other, "name": "Other", "url": "https://other.example.test", "monitor_type": "simple_web", "language": "", "web_server": ""},
                   "result": {"ok": False, "status": "down", "error_type": "unknown", "error": "legacy", "checked_at": datetime.now().isoformat()}}
        with self.db._conn:
            self.db._queue_event("legacy:batch-unknown", "down", payload, [rule["id"]])
        with patch.object(EmailNotifier, "_send", return_value=(True, "sent")) as send:
            EmailNotifier(self.db).process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 1)
        self.assertNotIn("Other", send.call_args.args[1])
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 1, "held": 1})

    def test_profile_and_rule_configuration_survives_restart_encrypted(self):
        profile = {"name": "VPN", "scope": "intranet_vpn", "enabled": True,
                   "targets": [{"url": "https://reference.example.test/", "method": "HEAD", "timeout_seconds": 2}]}
        self.db.save_network_profile(profile)
        rules = [{"id": "failure", "enabled": True, "priority": 1, "phase": "response",
                  "conditions": [{"source": "json", "path": "status", "operator": "equals", "value": "synthetic-private-value"}],
                  "result": {"error_type": "assertion_failed", "reason_code": "business_rule_match"}}]
        self.db.save_monitor({**self.db.get_monitor(self.id), "network_profile": "VPN", "diagnostic_rules_json": json.dumps(rules)})
        row = self.db._conn.execute("SELECT diagnostic_rules_enc FROM monitors WHERE id = ?", (self.id,)).fetchone()
        self.assertNotIn("synthetic-private-value", row[0])
        self.db.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db.get_monitor(self.id)["network_profile"], "VPN")
        self.assertEqual(json.loads(self.db.get_monitor(self.id)["diagnostic_rules_json"]), rules)
        self.assertEqual(self.db.get_network_profile("VPN")["scope"], "intranet_vpn")

    def test_unknown_profile_does_not_modify_monitor(self):
        before = self.db.get_monitor(self.id)
        with self.assertRaises(ValueError):
            self.db.save_monitor({**before, "name": "Changed", "network_profile": "Missing"})
        self.assertEqual(self.db.get_monitor(self.id)["name"], before["name"])

    def test_diagnosis_fields_are_persisted_with_logs_and_monitor(self):
        result = CheckResult(False, "down", status_code=503, error="HTTP error", error_type="http_error",
                             classification_state="confirmed", reason_code="http_503", phase="response",
                             evidence_source="http", evidence={"status_code": 503})
        self.db.record_check_result(self.id, result)
        log = self.db.list_logs()[0]
        monitor = self.db.get_monitor(self.id)
        self.assertEqual((log["classification_state"], log["reason_code"], log["phase"], log["evidence_source"]), ("confirmed", "http_503", "response", "http"))
        self.assertEqual(json.loads(log["evidence_json"]), {"status_code": 503})
        self.assertEqual(monitor["last_classification_state"], "confirmed")

    def test_checker_log_and_email_never_include_secrets_or_response_body(self):
        secrets = ["synthetic-userpass", "synthetic-query", "synthetic-token", "synthetic-header", "synthetic-response-only"]
        monitor = {**self.db.get_monitor(self.id), "monitor_type": "api",
                   "url": "https://user:synthetic-userpass@example.test/?token=synthetic-query",
                   "api_token": "synthetic-token", "headers_json": '{"X-Api-Key":"synthetic-header"}',
                   "expected_text": "required marker", "retry_count": 1}
        self.db.save_monitor(monitor)
        client = httpx.AsyncClient
        transport = httpx.MockTransport(lambda request: httpx.Response(200, text=" ".join(secrets)))
        with patch("sitesentry.core.checker.httpx.AsyncClient", side_effect=lambda **kwargs: client(transport=transport, **kwargs)):
            result = asyncio.run(SiteChecker().check_monitor(monitor))
        self.db.record_check_result(self.id, result)
        serialized = repr(result) + repr(self.db.list_logs()) + self.db._conn.execute("SELECT payload_json FROM notification_events").fetchone()[0]
        for secret in secrets:
            self.assertNotIn(secret, serialized)

    def test_rule_match_does_not_publish_condition_value(self):
        secret = "synthetic-rule-secret"
        rule = [{"id": "rule-1", "enabled": True, "priority": 1, "phase": "response",
                 "conditions": [{"source": "json", "path": "key", "operator": "equals", "value": secret}],
                 "result": {"error_type": "assertion_failed", "reason_code": "business_rule_match"}}]
        monitor = {**self.db.get_monitor(self.id), "monitor_type": "api", "retry_count": 1,
                   "diagnostic_rules_json": json.dumps(rule), "body": "not-json"}
        self.db.save_monitor(monitor)
        client = httpx.AsyncClient
        transport = httpx.MockTransport(lambda request: httpx.Response(200, json={"key": secret}))
        with patch("sitesentry.core.checker.httpx.AsyncClient", side_effect=lambda **kwargs: client(transport=transport, **kwargs)):
            result = asyncio.run(SiteChecker().check_monitor(monitor))
        self.db.record_check_result(self.id, result)
        payload = self.db._conn.execute("SELECT payload_json FROM notification_events").fetchone()[0]
        self.assertNotIn(secret, repr(result) + repr(self.db.list_logs()) + payload)
        self.assertEqual(result.evidence["rule_id"], "rule-1")

    def test_interrupted_smtp_attempt_becomes_uncertain_and_not_retried(self):
        self.record("timeout")
        delivery = self.db.pending_email_deliveries()[0]
        self.assertTrue(self.db.mark_delivery_attempt(delivery["delivery_id"]))
        self.db.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db.email_delivery_counts(), {"delivery_uncertain": 1})
        self.assertEqual(self.db.pending_email_deliveries(), [])
        self.db.retry_uncertain_delivery(delivery["delivery_id"])
        self.assertEqual(self.db.email_delivery_counts(), {"pending": 1})

    def test_connection_loss_during_send_is_uncertain(self):
        self.db.set_setting("smtp_host", "smtp.example.test")
        self.db.set_setting("smtp_to", "ops@example.test")
        self.db.set_setting("smtp_use_tls", "0")
        self.record("timeout")
        with patch("sitesentry.core.notifier.smtplib.SMTP") as factory:
            smtp = factory.return_value.__enter__.return_value
            smtp.sendmail.side_effect = smtplib.SMTPServerDisconnected("connection lost")
            notifier = EmailNotifier(self.db)
            notifier.process_pending()
            notifier.process_pending(now=datetime.now() + timedelta(minutes=2))
        self.assertEqual(smtp.sendmail.call_count, 1)
        self.assertEqual(self.db.email_delivery_counts(), {"delivery_uncertain": 1})

    def test_smtp_connection_failure_before_send_is_retryable(self):
        self.db.set_setting("smtp_host", "smtp.example.test")
        self.db.set_setting("smtp_to", "ops@example.test")
        self.record("timeout")
        with patch("sitesentry.core.notifier.smtplib.SMTP", side_effect=OSError("unreachable")):
            EmailNotifier(self.db).process_pending()
        self.assertEqual(self.db.email_delivery_counts(), {"failed": 1})
        row = self.db._conn.execute("SELECT error FROM email_deliveries").fetchone()
        self.assertIn("OSError", row[0])

    def test_disconnect_after_smtp_acceptance_does_not_mark_delivery_uncertain(self):
        self.db.set_setting("smtp_host", "smtp.example.test")
        self.db.set_setting("smtp_to", "ops@example.test")
        self.db.set_setting("smtp_use_tls", "0")
        self.record("timeout")
        with patch("sitesentry.core.notifier.smtplib.SMTP") as factory:
            factory.return_value.__enter__.return_value.sendmail.return_value = {}
            factory.return_value.__exit__.side_effect = smtplib.SMTPServerDisconnected("quit failed")
            EmailNotifier(self.db).process_pending()
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 1})

    def test_old_unknown_pending_event_is_held_on_upgrade(self):
        rule = self.db.list_email_rules()[0]
        monitor = self.db.get_monitor(self.id)
        payload = {"monitor": {"id": self.id, "name": monitor["name"], "url": monitor["url"], "monitor_type": "simple_web", "language": "", "web_server": ""},
                   "result": {"ok": False, "status": "down", "error_type": "unknown", "error": "legacy"}}
        with self.db._conn:
            self.db._queue_event("legacy:unknown", "down", payload, [rule["id"]])
        self.db.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db.email_delivery_counts(), {"held": 1})
        self.assertEqual(self.db.pending_email_deliveries(), [])

    def test_old_known_label_without_evidence_is_held_on_upgrade(self):
        rule = self.db.list_email_rules()[0]
        monitor = self.db.get_monitor(self.id)
        payload = {"monitor": {"id": self.id, "name": monitor["name"], "url": monitor["url"], "monitor_type": "simple_web", "language": "", "web_server": ""},
                   "result": {"ok": False, "status": "down", "error_type": "timeout", "error": "legacy"}}
        with self.db._conn:
            self.db._queue_event("legacy:unproven", "down", payload, [rule["id"]])
        self.db.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db.email_delivery_counts(), {"held": 1})
        self.assertEqual(self.db.pending_email_deliveries(), [])

    def test_v120_upgrade_preserves_secrets_and_rechecks_held_incident(self):
        self.db.save_monitor({**self.db.get_monitor(self.id), "auth_password": "fixture-only-password"})
        self.db.set_setting("smtp_password", "fixture-only-smtp", secret=True)
        self.record("timeout")
        event_id, payload_text = self.db._conn.execute("SELECT id, payload_json FROM notification_events").fetchone()
        payload = json.loads(payload_text)
        for key in ("classification_state", "reason_code", "phase", "evidence_source", "evidence"):
            payload["result"].pop(key, None)
        with self.db._conn:
            self.db._conn.execute("UPDATE notification_events SET payload_json = ? WHERE id = ?", (json.dumps(payload), event_id))
        self.db.close()
        connection = sqlite3.connect(self.temp.name + "/db.sqlite")
        for table, columns in {
            "monitors": ["network_profile", "diagnostic_rules_enc", "last_classification_state", "last_reason_code", "last_phase"],
            "logs": ["classification_state", "reason_code", "phase", "evidence_source", "evidence_json"],
            "email_incident_state": ["reason_code", "phase", "evidence_source", "evidence_json"],
        }.items():
            for column in columns:
                connection.execute(f"ALTER TABLE {table} DROP COLUMN {column}")
        connection.execute("DROP TABLE network_profiles")
        connection.execute("PRAGMA user_version = 0")
        connection.commit()
        connection.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db.get_monitor(self.id)["auth_password"], "fixture-only-password")
        self.assertEqual(self.db._conn.execute("PRAGMA user_version").fetchone()[0], SCHEMA_VERSION)
        self.assertEqual(self.db.get_setting("smtp_password"), "fixture-only-smtp")
        self.assertEqual(self.db.email_delivery_counts(), {"held": 1})
        self.assertEqual(self.db.list_logs()[0]["classification_state"], "legacy")
        self.record("timeout")
        self.assertEqual(self.db.email_delivery_counts(), {"held": 1, "pending": 1})
        self.db.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db.email_delivery_counts(), {"held": 1, "pending": 1})

    def test_partially_applied_migration_can_resume_without_duplicate_events(self):
        self.record("timeout")
        self.db.close()
        connection = sqlite3.connect(self.temp.name + "/db.sqlite")
        connection.execute("ALTER TABLE logs DROP COLUMN evidence_json")
        connection.execute("PRAGMA user_version = 0")
        connection.commit()
        connection.close()
        self.db = Database(self.temp.name + "/db.sqlite", SecurityManager(self.temp.name + "/key"))
        self.assertEqual(self.db._conn.execute("PRAGMA user_version").fetchone()[0], SCHEMA_VERSION)
        self.assertEqual(self.db._conn.execute("SELECT COUNT(*) FROM notification_events").fetchone()[0], 1)
        self.assertIn("evidence_json", self.db.list_logs()[0])

    def test_legacy_confirmed_http_delivery_is_sanitized_before_send(self):
        self.db.save_monitor({**self.db.get_monitor(self.id), "api_token": "legacy-secret", "url": "https://user:legacy-secret@example.test/?token=legacy-secret"})
        rule = self.db.list_email_rules()[0]
        payload = {"monitor": {"id": self.id, "name": "Unknown", "url": "https://user:legacy-secret@example.test/?token=legacy-secret", "monitor_type": "api", "language": "", "web_server": ""},
                   "result": {"ok": False, "status": "down", "error_type": "http_error", "status_code": 503, "error": "legacy-secret was exposed", "checked_at": datetime.now().isoformat()}}
        with self.db._conn:
            self.db._queue_event("legacy:http", "down", payload, [rule["id"]])
        with patch.object(EmailNotifier, "_send", return_value=(True, "sent")) as send:
            EmailNotifier(self.db).process_pending()
        self.assertEqual(send.call_count, 1)
        self.assertNotIn("legacy-secret", send.call_args.args[1])

    def test_csv_export_protects_formula_values(self):
        path = Path(self.temp.name) / "logs.csv"
        LogsView._write_csv([{"monitor_name": "=HYPERLINK(\"example\")", "error": "+cmd"}], str(path))
        data = path.read_text(encoding="utf-8-sig")
        self.assertIn("'=HYPERLINK", data)
        self.assertIn("'+cmd", data)

    def test_log_filters_and_exports_include_safe_diagnosis(self):
        self.db.record_check_result(self.id, CheckResult(False, "down", error_type="timeout", reason_code="read_timeout", phase="request", evidence_source="transport", evidence={"exception_kind": "ReadTimeout"}))
        self.record("unknown")
        filtered = self.db.list_logs(monitor_name="Unknown", classification_state="confirmed", phase="request")
        self.assertEqual(len(filtered), 1)
        self.assertEqual(filtered[0]["reason_code"], "read_timeout")
        path = Path(self.temp.name) / "logs.xlsx"
        LogsView._write_excel(filtered, str(path))
        book = load_workbook(path)
        try:
            headers = [cell.value for cell in book.active[1]]
            self.assertIn("classification_state", headers)
            self.assertIn("evidence_json", headers)
        finally:
            book.close()
    def test_report_breakdown_preserves_unknown_and_reason_counts(self):
        self.db.record_check_result(self.id, CheckResult(False, "down", error_type="timeout", reason_code="read_timeout", phase="request", evidence_source="transport", evidence={"exception_kind": "ReadTimeout"}))
        self.record("unknown")
        now = datetime.now()
        report = build_report(self.db, now - timedelta(minutes=1), now + timedelta(minutes=1), now=now)
        row = report["rows"][0]
        self.assertEqual(row["down_count"], 2)
        self.assertEqual(sum(item["count"] for item in row["diagnostic_counts"]), 2)
        self.assertIn("read_timeout", report["incidents"][0]["reason_codes"])
        self.assertEqual(set(report["incidents"][0]["classification_states"]), {"confirmed", "undetermined"})
        path = Path(self.temp.name) / "statistics.xlsx"
        export_report(report, path)
        book = load_workbook(path)
        try:
            headers = [cell.value for cell in book["Error breakdown"][1]]
            self.assertIn("Classification", headers)
            self.assertIn("Reason code", headers)
            self.assertEqual(book["Error breakdown"].max_row, 3)
            self.assertIn("read_timeout", book["Incidents"]["J2"].value)
        finally:
            book.close()
