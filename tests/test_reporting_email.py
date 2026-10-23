from __future__ import annotations

from datetime import datetime, timedelta
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from sitesentry.core.checker import CheckResult
from sitesentry.core.database import Database
from sitesentry.core.email_rules import normalize_rule, parse_recipients, render_template
from sitesentry.core.excel_importer import MONITOR_COLUMNS, create_monitor_template, import_monitors_from_excel
from sitesentry.core.notifier import EmailNotifier
from sitesentry.core.reports import build_report, export_report, format_duration
from sitesentry.core.scheduler import MonitorScheduler
from sitesentry.core.security import SecurityManager


class DatabaseCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.security = SecurityManager(self.root / "key")
        self.db = Database(self.root / "test.db", self.security)
        self.db.set_setting("smtp_enabled", "1")
        self.db.set_setting("smtp_host", "smtp.example.test")
        self.db.set_setting("smtp_to", "first@example.test; second@example.test")
        self.db.set_setting("smtp_use_tls", "0")
        self.monitor_id = self.db.save_monitor({"name": "Portal", "url": "https://example.test", "language": "PHP", "web_server": "Apache"})

    def tearDown(self) -> None:
        self.db.close()
        self.temp.cleanup()

    def result(self, ok: bool, minute: int = 0) -> CheckResult:
        return CheckResult(ok, "up" if ok else "down", 123, 200 if ok else 503, checked_at=f"2026-10-20T10:{minute:02d}:00")

    def log(self, stamp: str, status: str, monitor_id: int | None = None) -> None:
        self.db.add_log(monitor_id or self.monitor_id, "Portal", status, status == "up", 123, 200, checked_at=stamp)

    def retry_ready(self) -> None:
        with self.db._lock, self.db._conn:
            self.db._conn.execute("UPDATE email_deliveries SET last_attempt_at = NULL WHERE status != 'sent'")


class EmailTests(DatabaseCase):
    def test_transitions_are_persistent_and_ignore_cooldown(self) -> None:
        for minute, ok in enumerate([False, False, True, True, False, True]):
            self.db.record_check_result(self.monitor_id, self.result(ok, minute))
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending()
        self.assertEqual(send.call_count, 4)
        self.assertEqual([call.args[0] for call in send.call_args_list], ["[SiteSentry] DOWN: Portal", "[SiteSentry] UP: Portal", "[SiteSentry] DOWN: Portal", "[SiteSentry] UP: Portal"])
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 4})
        self.db.close()
        self.db = Database(self.root / "test.db", self.security)
        self.db.record_check_result(self.monitor_id, self.result(False, 7))
        self.db.record_check_result(self.monitor_id, self.result(False, 8))
        self.assertEqual(len(self.db.pending_email_deliveries()), 1)
        self.assertEqual(len(self.db.list_logs()), 8)

    def test_disabled_smtp_and_disabled_monitor_do_not_queue(self) -> None:
        self.db.set_setting("smtp_enabled", "0")
        self.db.record_check_result(self.monitor_id, self.result(False))
        self.assertEqual(self.db.email_delivery_counts(), {})
        self.db.set_monitor_enabled(self.monitor_id, False)
        self.assertFalse(self.db.record_check_result(self.monitor_id, self.result(True)))
        self.assertEqual(len(self.db.list_logs()), 1)

    def test_restart_while_down_does_not_create_another_alert(self) -> None:
        self.db.record_check_result(self.monitor_id, self.result(False))
        self.db.close()
        self.db = Database(self.root / "test.db", self.security)
        self.db.record_check_result(self.monitor_id, self.result(False, 1))
        self.db.record_check_result(self.monitor_id, self.result(True, 2))
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending()
        self.assertEqual(send.call_count, 2)

    def test_background_worker_wakes_on_new_event(self) -> None:
        accepted = threading.Event()
        notifier = EmailNotifier(self.db)

        def send(*args, **kwargs):
            accepted.set()
            return True, "sent"

        with patch.object(notifier, "_send", side_effect=send):
            notifier.start()
            try:
                self.db.record_check_result(self.monitor_id, self.result(False))
                notifier.wake()
                self.assertTrue(accepted.wait(5))
            finally:
                notifier.stop()
                notifier._worker.join(5)
        self.assertFalse(notifier._worker.is_alive())

    def test_monitor_name_cannot_inject_subject_headers(self) -> None:
        monitor = self.db.get_monitor(self.monitor_id)
        self.db.save_monitor({**monitor, "name": "Portal\nBcc: attacker@example.test"})
        self.db.record_check_result(self.monitor_id, self.result(False))
        subject, _ = EmailNotifier(self.db)._rule_message(self.db.pending_email_deliveries()[0])
        self.assertNotIn("\n", subject)

    def test_summary_contains_daily_dashboard_and_custom_values(self) -> None:
        rule_id = self.db.save_email_rule({"name": "Daily", "rule_type": "daily_summary", "weekdays": "mon", "body_template": "{date}: {total} / {down_count} / {downtime}"})
        self.log("2026-10-20T01:00:00", "down")
        report = build_report(self.db, datetime(2026, 10, 20), datetime(2026, 10, 20, 2))
        self.db.queue_summary(rule_id, "qa", report)
        subject, body = EmailNotifier(self.db)._rule_message(self.db.pending_email_deliveries()[0])
        self.assertEqual(subject, "[SiteSentry] Daily summary 2026-10-20")
        self.assertIn("2026-10-20: 1 / 1 / 01:00:00", body)
        self.assertIn("Total: 1 | UP: 0 | DOWN: 1 | Unknown: 0", body)

    def test_recovery_only_and_multiple_rules(self) -> None:
        recovery = self.db.save_email_rule({"name": "Recovery", "rule_type": "recovery_only", "recipients": "ops@example.test"})
        self.db.record_check_result(self.monitor_id, self.result(False))
        self.db.record_check_result(self.monitor_id, self.result(False, 1))
        self.db.record_check_result(self.monitor_id, self.result(True, 2))
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending()
        self.assertEqual(send.call_count, 3)
        self.assertEqual(sum(call.kwargs["recipient_override"] == "ops@example.test" for call in send.call_args_list), 1)
        self.db.delete_email_rule(recovery)

    def test_failure_retries_in_order_without_duplicate_down(self) -> None:
        for minute, ok in enumerate([False, False, True]):
            self.db.record_check_result(self.monitor_id, self.result(ok, minute))
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(False, "offline")) as send:
            notifier.process_pending()
            notifier.process_pending()
        self.assertEqual(send.call_count, 1)
        self.assertEqual(self.db.pending_email_deliveries(), [])
        self.retry_ready()
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending()
        self.assertEqual(send.call_count, 2)
        self.assertIn("DOWN", send.call_args_list[0].args[0])
        self.assertIn("UP", send.call_args_list[1].args[0])

    def test_partial_refusal_only_retries_refused_recipient(self) -> None:
        self.db.record_check_result(self.monitor_id, self.result(False))
        notifier = EmailNotifier(self.db)
        with patch("sitesentry.core.notifier.smtplib.SMTP") as smtp_class:
            smtp = smtp_class.return_value.__enter__.return_value
            smtp.sendmail.side_effect = [{"second@example.test": (450, b"try later")}, {}]
            notifier.process_pending()
            self.retry_ready()
            notifier.process_pending()
            self.assertEqual(smtp.sendmail.call_args_list[0].args[1], ["first@example.test", "second@example.test"])
            self.assertEqual(smtp.sendmail.call_args_list[1].args[1], ["second@example.test"])
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 1})

    def test_custom_body_and_subject_escape_html(self) -> None:
        rule = self.db.list_email_rules()[0]
        self.db.save_email_rule({**rule, "subject_template": "{monitor_name}: {status}", "body_template": "Hello <team>\n{error}", "include_details": False})
        self.db.record_check_result(self.monitor_id, self.result(False))
        delivery = self.db.pending_email_deliveries()[0]
        subject, body = EmailNotifier(self.db)._rule_message(delivery)
        self.assertEqual(subject, "Portal: DOWN")
        self.assertIn("&lt;team&gt;<br>", body)
        self.assertNotIn("Incident detected", body)
        self.assertNotIn("auth_password", delivery["payload_json"])

    def test_rule_validation_and_recipient_parsing(self) -> None:
        self.assertEqual(parse_recipients("A <A@example.test>; b@example.test\na@example.test"), ["A@example.test", "b@example.test"])
        for payload in [{"name": ""}, {"name": "x", "rule_type": "bad"}, {"name": "x", "rule_type": "daily_summary", "weekdays": ""}, {"name": "x", "subject_template": "{obj.__class__}"}, {"name": "x", "subject_template": "bad\nheader"}, {"name": "x", "include_details": False}, {"name": "x", "recipients": "invalid"}]:
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                normalize_rule(payload)
        self.assertEqual(render_template("{{test}} {status}", {"status": "UP"}), "{test} UP")

    def test_default_rule_not_recreated_after_deletion(self) -> None:
        for rule in self.db.list_email_rules():
            self.db.delete_email_rule(rule["id"])
        self.db.init_schema()
        self.assertEqual(self.db.list_email_rules(), [])


class ReportTests(DatabaseCase):
    def test_down_count_and_incident_average_are_different(self) -> None:
        for stamp, state in [("10:00", "down"), ("10:05", "down"), ("10:10", "up"), ("10:20", "down"), ("10:40", "up")]:
            self.log(f"2026-10-20T{stamp}:00", state)
        report = build_report(self.db, datetime(2026, 10, 20), datetime(2026, 10, 21))
        row = report["rows"][0]
        self.assertEqual((row["down_count"], row["incident_count"], row["downtime_seconds"], row["average_down_seconds"]), (3, 2, 1800, 900))
        self.assertEqual(report["up"], 1)

    def test_clips_carried_and_open_outages_to_period(self) -> None:
        self.log("2026-10-19T23:00:00", "down")
        self.log("2026-10-20T01:00:00", "up")
        self.log("2026-10-20T02:00:00", "down")
        self.log("2026-10-20T03:00:00", "down")
        report = build_report(self.db, datetime(2026, 10, 20), datetime(2026, 10, 20, 23, 59, 59), now=datetime(2026, 10, 20, 4))
        row = report["rows"][0]
        self.assertEqual((row["down_count"], row["incident_count"], row["downtime_seconds"], row["open_incidents"]), (2, 2, 10800, 1))
        self.assertTrue(report["incidents"][0]["carried_in"])
        self.assertIsNone(report["incidents"][1]["recovered_at"])

    def test_recovery_exactly_at_start_not_an_incident(self) -> None:
        self.log("2026-10-19T23:00:00", "down")
        self.log("2026-10-20T00:00:00", "up")
        report = build_report(self.db, datetime(2026, 10, 20), datetime(2026, 10, 21))
        self.assertEqual(report["rows"][0]["incident_count"], 0)

    def test_more_than_500_logs_and_deleted_monitor(self) -> None:
        for minute in range(610):
            self.log((datetime(2026, 10, 20) + timedelta(minutes=minute)).isoformat(), "down")
        self.log("2026-10-20T11:00:00", "up")
        self.db.delete_monitor(self.monitor_id)
        report = build_report(self.db, datetime(2026, 10, 20), datetime(2026, 10, 21))
        self.assertEqual(report["down_count"], 610)
        self.assertEqual(report["downtime_seconds"], 39600)
        self.assertEqual(report["total"], 0)
        self.assertEqual(report["rows"][0]["monitor_name"], "Portal")

    def test_invalid_and_future_period(self) -> None:
        with self.assertRaises(ValueError):
            build_report(self.db, datetime(2026, 10, 21), datetime(2026, 10, 20))
        report = build_report(self.db, datetime(2026, 10, 21), datetime(2026, 10, 22), now=datetime(2026, 10, 20))
        self.assertEqual(report["downtime_seconds"], 0)
        self.assertEqual(report["unknown"], 1)

    def test_export_has_typed_dates_durations_and_safe_text(self) -> None:
        monitor = self.db.get_monitor(self.monitor_id)
        self.db.save_monitor({**monitor, "name": "=HYPERLINK(\"bad\")"})
        self.log("2026-10-19T22:00:00", "down")
        self.log("2026-10-21T02:00:00", "up")
        report = build_report(self.db, datetime(2026, 10, 19), datetime(2026, 10, 22))
        path = self.root / "report.xlsx"
        export_report(report, path)
        workbook = load_workbook(path)
        sheet = workbook["Statistics"]
        self.assertEqual(sheet["A6"].data_type, "s")
        self.assertEqual(sheet["D6"].number_format, "[h]:mm:ss")
        self.assertEqual(sheet["D6"].value, timedelta(hours=28))
        self.assertIsInstance(sheet["B2"].value, datetime)
        self.assertEqual(workbook["Incidents"]["E2"].value, timedelta(hours=28))
        workbook.close()
        self.assertEqual(format_duration(100000), "27:46:40")


class SchedulerTests(DatabaseCase):
    def test_multiple_schedules_update_remove_and_deduplicate(self) -> None:
        scheduler = MonitorScheduler(self.db, MagicMock(), MagicMock())
        ids = [self.db.save_email_rule({"name": f"Digest {hour}", "rule_type": "daily_summary", "hour": hour, "weekdays": "mon,tue,wed,thu,fri", "report_period": "yesterday"}) for hour in [8, 18]]
        scheduler.sync_email_jobs()
        scheduler.sync_email_jobs()
        self.assertEqual(len(scheduler._scheduler.get_jobs()), 2)
        self.assertTrue(scheduler.queue_daily_summary(ids[0], datetime(2026, 10, 20, 8)))
        self.assertFalse(scheduler.queue_daily_summary(ids[0], datetime(2026, 10, 20, 8)))
        self.assertTrue(scheduler.queue_daily_summary(ids[1], datetime(2026, 10, 20, 18)))
        payload = json.loads(self.db.pending_email_deliveries()[0]["payload_json"])
        self.assertEqual(payload["report"]["start_at"], "2026-10-19T00:00:00")
        self.assertEqual(payload["report"]["end_at"], "2026-10-19T23:59:59")
        self.db.delete_email_rule(ids[0])
        scheduler.sync_email_jobs()
        self.assertEqual(len(scheduler._scheduler.get_jobs()), 1)

    def test_monitor_cannot_overlap_manual_and_scheduled_checks(self) -> None:
        entered, release = threading.Event(), threading.Event()
        checker = MagicMock()

        async def check(monitor: dict) -> CheckResult:
            entered.set()
            release.wait(5)
            return self.result(False)

        checker.check_monitor.side_effect = check
        scheduler = MonitorScheduler(self.db, checker, MagicMock())
        first = threading.Thread(target=scheduler._run_monitor_job, args=(self.monitor_id,))
        first.start()
        self.assertTrue(entered.wait(5))
        scheduler._run_monitor_job(self.monitor_id)
        release.set()
        first.join(5)
        self.assertFalse(first.is_alive())
        self.assertEqual(checker.check_monitor.call_count, 1)
        self.assertEqual(len(self.db.list_logs()), 1)


class ImportTests(DatabaseCase):
    def test_template_and_import_metadata_all_types(self) -> None:
        path = create_monitor_template(self.root / "template.xlsx")
        workbook = load_workbook(path)
        self.assertEqual([cell.value for cell in workbook["Monitors"][1]], MONITOR_COLUMNS)
        self.assertEqual(workbook["Monitors"]["AC1"].value, "language")
        self.assertEqual(workbook["Monitors"]["AD1"].value, "web_server")
        workbook.close()
        result = import_monitors_from_excel(self.db, path)
        self.assertEqual(result.created, 3)
        self.assertFalse(result.errors)
        imported = [row for row in self.db.list_monitors() if row["id"] != self.monitor_id]
        self.assertEqual({row["language"] for row in imported}, {"PHP", "Python", "ASP.NET"})

    def test_legacy_template_preserves_metadata(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["name", "monitor_type", "url"])
        sheet.append(["Portal", "simple_web", "https://example.test"])
        path = self.root / "old.xlsx"
        workbook.save(path)
        result = import_monitors_from_excel(self.db, path)
        self.assertEqual(result.updated, 1)
        monitor = self.db.get_monitor(self.monitor_id)
        self.assertEqual((monitor["language"], monitor["web_server"]), ("PHP", "Apache"))

    def test_schema_upgrade_preserves_old_data_and_secrets(self) -> None:
        self.db.set_setting("smtp_password", "secret", secret=True)
        self.db.close()
        connection = sqlite3.connect(self.root / "test.db")
        connection.executescript("ALTER TABLE monitors DROP COLUMN language; ALTER TABLE monitors DROP COLUMN web_server;")
        connection.close()
        self.db = Database(self.root / "test.db", self.security)
        self.assertEqual(self.db.get_monitor(self.monitor_id)["language"], "")
        self.assertEqual(self.db.get_setting("smtp_password"), "secret")
        self.assertEqual(len(self.db.list_email_rules()), 1)


if __name__ == "__main__":
    unittest.main()
