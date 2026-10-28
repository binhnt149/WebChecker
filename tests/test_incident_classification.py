import asyncio
import errno
import json
import tempfile
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path
import unittest
import random
from unittest.mock import patch

import httpx
from openpyxl import load_workbook

from sitesentry.core.checker import CheckResult, SiteChecker
from sitesentry.core.database import Database
from sitesentry.core.email_rules import normalize_rule
from sitesentry.core.notifier import EmailNotifier
from sitesentry.core.security import SecurityManager
from sitesentry.core.reports import build_report
from sitesentry.core.report_export import create_report_template, export_report


class ClassificationTests(unittest.TestCase):
    def run_check(self, handler, monitor=None):
        real_client = httpx.AsyncClient
        transport = httpx.MockTransport(handler)
        with patch("sitesentry.core.checker.httpx.AsyncClient", side_effect=lambda **kwargs: real_client(transport=transport, **kwargs)):
            return asyncio.run(SiteChecker().check_monitor({"url": "https://example.test", "retry_count": 1, **(monitor or {})}))

    def test_timeout_is_not_incorrect_password(self):
        def handler(request):
            raise httpx.ReadTimeout("read timed out", request=request)
        result = self.run_check(handler)
        self.assertEqual(result.error_type, "timeout")

    def test_auth_transport_failures_keep_network_category(self):
        def handler(request):
            raise httpx.ConnectError("connection refused", request=request)
        with patch("sitesentry.core.diagnostics.has_local_route", return_value=True):
            result = self.run_check(handler, {"monitor_type": "auth_web", "auth_login_url": "https://example.test/login"})
        self.assertEqual(result.error_type, "target_unreachable")
        self.assertIn("Login page failed", result.error)

    def test_network_down_requires_local_evidence(self):
        def handler(request):
            raise httpx.ConnectError("no route", request=request) from OSError(errno.ENETUNREACH, "Network unreachable")
        self.assertEqual(self.run_check(handler).error_type, "network_offline")

    def test_server_401_and_403_are_distinct(self):
        self.assertEqual(self.run_check(lambda req: httpx.Response(401)).error_type, "auth_credentials")
        self.assertEqual(self.run_check(lambda req: httpx.Response(403)).error_type, "http_error")

    def test_missing_text_is_an_assertion(self):
        result = self.run_check(lambda req: httpx.Response(200, text="Hello"), {"expected_text": "Dashboard"})
        self.assertEqual(result.error_type, "assertion_failed")

    def test_auth_login_rejection_is_not_retried(self):
        requests = []
        form = '<form><input name="username"><input name="password" type="password"></form>'
        def handler(request):
            requests.append(request.method)
            return httpx.Response(200, text=form + ("Invalid username or password" if request.method == "POST" else ""))
        result = self.run_check(handler, {"monitor_type": "auth_web", "auth_login_url": "https://example.test/login", "retry_count": 3})
        self.assertEqual(result.error_type, "auth_credentials")
        self.assertEqual(requests, ["GET", "POST"])

    def test_login_page_without_explicit_rejection_is_not_bad_password(self):
        form = '<form><input name="username"><input name="password" type="password"></form>'
        result = self.run_check(lambda req: httpx.Response(200, text=form), {"monitor_type": "auth_web", "auth_login_url": "https://example.test"})
        self.assertEqual(result.error_type, "auth_verification")

    def test_json_failure_has_category(self):
        result = self.run_check(lambda req: httpx.Response(200, json={"status": "bad"}), {"monitor_type": "api", "expected_json_path": "status", "expected_json_value": "ok"})
        self.assertEqual(result.error_type, "assertion_failed")

    def test_successful_login_with_search_form_and_help_text_is_not_rejected(self):
        login = '<form><input name="username"><input name="password" type="password"></form>'
        def handler(request):
            if request.method == "POST":
                return httpx.Response(200, text='<form><input name="search"></form>If authentication failed, contact support')
            return httpx.Response(200, text=login if request.url.path == "/login" else "Dashboard")
        result = self.run_check(handler, {"monitor_type": "auth_web", "auth_login_url": "https://example.test/login", "auth_verify_url": "https://example.test/dashboard", "auth_verify_text": "Dashboard"})
        self.assertTrue(result.ok, result.error)

    def test_hidden_login_error_on_503_remains_server_error(self):
        form = '<form><input name="username"><input name="password" type="password"></form>'
        def handler(request):
            return httpx.Response(503 if request.method == "POST" else 200, text=form + '<div hidden>Invalid credentials</div>')
        result = self.run_check(handler, {"monitor_type": "auth_web", "auth_login_url": "https://example.test/login"})
        self.assertEqual(result.error_type, "http_error")

    def test_hidden_login_error_on_success_is_not_a_credential_rejection(self):
        form = '<form><input name="username"><input name="password" type="password"></form>'
        result = self.run_check(lambda req: httpx.Response(200, text=form + '<div style="display: none">Invalid credentials</div>'), {"monitor_type": "auth_web", "auth_login_url": "https://example.test"})
        self.assertEqual(result.error_type, "auth_verification")

    def test_inline_formatted_credential_rejection_is_detected(self):
        form = '<form><input name="username"><input name="password" type="password"></form>'
        def handler(request):
            return httpx.Response(200, text=form + ('<p>Invalid <strong>username or password</strong>.</p>' if request.method == "POST" else ""))
        result = self.run_check(handler, {"monitor_type": "auth_web", "auth_login_url": "https://example.test/login"})
        self.assertEqual(result.error_type, "auth_credentials")


class IncidentPolicyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.db = Database(root / "db", SecurityManager(root / "key"))
        self.db.set_setting("smtp_enabled", "1")
        self.rule = self.db.list_email_rules()[0]
        self.db.save_email_rule({**self.rule, "batch_window_seconds": 15, "excluded_error_types": ["auth_credentials"]})
        self.monitors = [self.db.save_monitor({"name": f"Monitor {index}", "url": f"https://example.test/{index}", "monitor_type": kind}) for index, kind in enumerate(["simple_web", "auth_web", "api"])]

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def record(self, monitor_id, category="timeout", ok=False):
        result = CheckResult(ok, "up" if ok else "down", error_type="" if ok else category, error="test only")
        self.db.record_check_result(monitor_id, result)

    def test_exclusion_preserves_log_but_suppresses_down_and_recovery(self):
        self.record(self.monitors[0], "auth_credentials")
        self.record(self.monitors[0], ok=True)
        self.assertEqual(len(self.db.list_logs()), 2)
        self.assertEqual(self.db.list_logs()[-1]["error_type"], "auth_credentials")
        self.assertEqual(self.db.email_delivery_counts(), {})

    def test_excluded_error_becoming_eligible_while_still_down_alerts_once(self):
        self.record(self.monitors[0], "auth_credentials")
        self.record(self.monitors[0], "timeout")
        self.record(self.monitors[0], "timeout")
        self.assertEqual(self.db.email_delivery_counts(), {"pending": 1})

    def test_three_monitor_types_become_one_email_after_window(self):
        for monitor_id in self.monitors:
            self.record(monitor_id)
        self.assertEqual(self.db.pending_email_deliveries(), [])
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 1)
        for index in range(3):
            self.assertIn(f"Monitor {index}", send.call_args.args[1])
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 3})
        with patch.object(notifier, "_send") as send:
            notifier.process_pending(now=datetime.now() + timedelta(minutes=2))
        send.assert_not_called()

    def test_batch_retry_and_recovery_order(self):
        for monitor_id in self.monitors:
            self.record(monitor_id)
        self.record(self.monitors[0], ok=True)
        notifier = EmailNotifier(self.db)
        now = datetime.now() + timedelta(seconds=20)
        with patch.object(notifier, "_send", return_value=(False, "offline")) as send:
            notifier.process_pending(now=now)
        self.assertEqual(send.call_count, 1)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=now + timedelta(minutes=2))
        self.assertEqual(send.call_count, 2)
        self.assertIn("DOWN", send.call_args_list[0].args[0])
        self.assertIn("UP", send.call_args_list[1].args[0])

    def test_validation_rejects_unknown_errors_and_unbounded_windows(self):
        for data in [{"excluded_error_types": ["typo"]}, {"batch_window_seconds": -1}, {"batch_window_seconds": 301}]:
            with self.subTest(data=data), self.assertRaises(ValueError):
                normalize_rule({"name": "rule", **data})

    def test_exclusions_changed_during_grouping_apply_before_send(self):
        self.record(self.monitors[0], "timeout")
        self.record(self.monitors[1], "http_error")
        self.record(self.monitors[0], ok=True)
        self.db.save_email_rule({**self.db.get_email_rule(self.rule["id"]), "excluded_error_types": ["timeout"]})
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 1)
        self.assertNotIn("Monitor 0", send.call_args.args[1])
        self.assertIn("Monitor 1", send.call_args.args[1])
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 1, "suppressed": 2})

    def test_batch_partial_delivery_retry_does_not_resend_accepted_addresses(self):
        self.db.set_setting("smtp_host", "smtp.example.test")
        self.db.set_setting("smtp_to", "one@example.test,two@example.test")
        self.db.set_setting("smtp_use_tls", "0")
        for monitor_id in self.monitors:
            self.record(monitor_id)
        notifier = EmailNotifier(self.db)
        with patch("sitesentry.core.notifier.smtplib.SMTP") as smtp_class:
            smtp = smtp_class.return_value.__enter__.return_value
            smtp.sendmail.side_effect = [{"two@example.test": (450, b"later")}, {}]
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
            notifier.process_pending(now=datetime.now() + timedelta(minutes=2))
        self.assertEqual(smtp.sendmail.call_count, 2)
        self.assertEqual(smtp.sendmail.call_args_list[-1].args[1], ["two@example.test"])
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 3})

    def test_report_breakdown_keeps_all_excluded_errors(self):
        self.record(self.monitors[0], "auth_credentials")
        self.record(self.monitors[1], "timeout")
        now = datetime.now()
        report = build_report(self.db, now.replace(hour=0, minute=0, second=0, microsecond=0), now + timedelta(minutes=1))
        self.assertEqual(report["error_counts"], {"auth_credentials": 1, "timeout": 1})
        self.assertEqual(report["rows"][0]["last_error_type"], "auth_credentials")
        self.assertEqual(report["incidents"][0]["error_types"], ["auth_credentials"])

    def test_export_contains_diagnostics_and_usable_empty_template(self):
        self.record(self.monitors[0], "timeout")
        now = datetime.now()
        report = build_report(self.db, now.replace(hour=0, minute=0, second=0, microsecond=0), now)
        path = Path(self.temp.name) / "report.xlsx"
        export_report(report, path)
        workbook = load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Statistics", "Incidents", "Error breakdown"])
        self.assertEqual(workbook["Error breakdown"]["B2"].value, "timeout")
        self.assertEqual(workbook["Error breakdown"]["D2"].value, 1)
        self.assertEqual(workbook["Statistics"].freeze_panes, "B6")
        self.assertFalse(workbook["Statistics"].sheet_view.showGridLines)
        self.assertEqual(workbook["Statistics"]["D6"].number_format, "[h]:mm:ss")
        workbook.close()
        create_report_template(path)
        workbook = load_workbook(path)
        self.assertEqual(workbook["Statistics"].max_row, 5)
        self.assertEqual(workbook["Incidents"].max_row, 1)
        self.assertIsInstance(workbook["Statistics"]["B2"].value, datetime)
        workbook.close()

    def test_daily_summary_filters_only_current_excluded_down_rows(self):
        from sitesentry.core.reports import filter_email_summary
        self.record(self.monitors[0], "auth_credentials")
        self.record(self.monitors[1], "timeout")
        now = datetime.now()
        report = build_report(self.db, now.replace(hour=0, minute=0, second=0, microsecond=0), now)
        filtered = filter_email_summary(report, {"auth_credentials"})
        self.assertEqual(filtered["excluded_monitor_count"], 1)
        self.assertEqual(filtered["down"], 1)
        self.assertEqual(len(report["rows"]), 3)

    def test_fast_repeated_incident_is_not_counted_as_two_monitors(self):
        self.record(self.monitors[0])
        self.record(self.monitors[0], ok=True)
        self.record(self.monitors[0])
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 3)
        self.assertTrue(all("2 monitors" not in call.args[0] for call in send.call_args_list))

    def test_filter_change_between_fetch_and_claim_never_sends_excluded_alert(self):
        self.record(self.monitors[0])
        original = self.db.mark_delivery_attempt
        def claim(delivery_id):
            self.db.save_email_rule({**self.db.get_email_rule(self.rule["id"]), "excluded_error_types": ["timeout"]})
            return original(delivery_id)
        notifier = EmailNotifier(self.db)
        with patch.object(self.db, "mark_delivery_attempt", side_effect=claim), patch.object(notifier, "_send") as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        send.assert_not_called()

    def test_excluding_older_incident_does_not_clear_newer_incident(self):
        self.record(self.monitors[0], "timeout")
        self.record(self.monitors[0], ok=True)
        self.record(self.monitors[0], "http_error")
        self.db.save_email_rule({**self.db.get_email_rule(self.rule["id"]), "excluded_error_types": ["timeout"]})
        self.record(self.monitors[0], "http_error")
        self.record(self.monitors[0], ok=True)
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 2)
        self.assertIn("DOWN", send.call_args_list[0].args[0])
        self.assertIn("UP", send.call_args_list[1].args[0])

    def test_interleaved_incidents_preserve_recovery_order_for_each_monitor(self):
        a, b = self.monitors[:2]
        for monitor_id, ok in [(a, False), (b, False), (b, True), (b, False), (a, True), (a, False)]:
            self.record(monitor_id, ok=ok)
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 5)
        self.assertIn("UP: Monitor 0", send.call_args_list[3].args[0])
        self.assertIn("DOWN: Monitor 0", send.call_args_list[4].args[0])

    def test_new_filter_suppresses_unattempted_recovery_only_event(self):
        self.db.delete_email_rule(self.rule["id"])
        rule_id = self.db.save_email_rule({"name": "Recovery", "rule_type": "recovery_only"})
        self.record(self.monitors[0])
        self.record(self.monitors[0], ok=True)
        self.db.save_email_rule({**self.db.get_email_rule(rule_id), "excluded_error_types": ["timeout"]})
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send") as send:
            notifier.process_pending()
        send.assert_not_called()

    def test_unrelated_recovery_does_not_split_concurrent_down_alerts(self):
        a, b, c = self.monitors
        self.record(c)
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")):
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.record(a)
        self.record(c, ok=True)
        self.record(b)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 2)
        self.assertIn("Monitor 0", send.call_args_list[0].args[1])
        self.assertIn("Monitor 1", send.call_args_list[0].args[1])

    def test_rule_type_change_keeps_original_recovery_exclusion_semantics(self):
        self.db.delete_email_rule(self.rule["id"])
        rule_id = self.db.save_email_rule({"name": "Recovery", "rule_type": "recovery_only"})
        self.record(self.monitors[0])
        self.record(self.monitors[0], ok=True)
        self.db.save_email_rule({**self.db.get_email_rule(rule_id), "rule_type": "incident", "excluded_error_types": ["timeout"]})
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send") as send:
            notifier.process_pending()
        send.assert_not_called()
        self.assertEqual(self.db.email_delivery_counts(), {"suppressed": 1})

    def test_random_interleaved_incidents_keep_per_monitor_event_order(self):
        rng = random.Random(784)
        states = {mid: None for mid in self.monitors}
        expected = {mid: [] for mid in self.monitors}
        for _ in range(80):
            mid, ok = rng.choice(self.monitors), rng.choice([True, False])
            if (not ok and states[mid] is not False) or (ok and states[mid] is False):
                expected[mid].append("up" if ok else "down")
            self.record(mid, ok=ok)
            states[mid] = ok
        seen = {mid: [] for mid in self.monitors}
        notifier = EmailNotifier(self.db)
        render = notifier._rule_message
        def capture(delivery):
            members = self.db.email_batch_members(delivery["batch_id"]) if delivery.get("batch_id") else [delivery]
            for member in members:
                payload = json.loads(member["payload_json"])
                seen[payload["monitor"]["id"]].append(payload["result"]["status"])
            return render(delivery)
        with patch.object(notifier, "_rule_message", side_effect=capture), patch.object(notifier, "_send", return_value=(True, "sent")):
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(seen, expected)

    def test_upgrade_legacy_queue_preserves_secrets_and_does_not_repeat_down(self):
        monitor = self.db.get_monitor(self.monitors[0])
        self.db.save_monitor({**monitor, "auth_password": "temporary-test-secret"})
        self.db.set_setting("smtp_password", "smtp-test-secret", secret=True)
        self.record(self.monitors[0], "http_error")
        root = Path(self.temp.name)
        self.db.close()
        connection = sqlite3.connect(root / "db")
        connection.executescript("""
            ALTER TABLE email_deliveries DROP COLUMN batch_id;
            ALTER TABLE email_deliveries DROP COLUMN rule_type_at_creation;
            DROP TABLE email_batches;
            DROP TABLE email_incident_state;
            ALTER TABLE email_rules DROP COLUMN excluded_error_types_json;
            ALTER TABLE email_rules DROP COLUMN batch_window_seconds;
            ALTER TABLE monitors DROP COLUMN last_error_type;
            ALTER TABLE logs DROP COLUMN error_type;
            DELETE FROM settings WHERE key = 'incident_policy_initialized';
        """)
        connection.close()
        self.db = Database(root / "db", SecurityManager(root / "key"))
        self.assertEqual(self.db.get_setting("smtp_password"), "smtp-test-secret")
        self.assertEqual(self.db.get_monitor(self.monitors[0])["auth_password"], "temporary-test-secret")
        self.record(self.monitors[0], "http_error")
        self.record(self.monitors[0], ok=True)
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send", return_value=(True, "sent")) as send:
            notifier.process_pending(now=datetime.now() + timedelta(seconds=20))
        self.assertEqual(send.call_count, 2)
        self.assertEqual(self.db.email_delivery_counts(), {"sent": 2})

    def test_unknown_legacy_recovery_provenance_is_filtered_conservatively(self):
        self.db.delete_email_rule(self.rule["id"])
        rule_id = self.db.save_email_rule({"name": "Recovery", "rule_type": "recovery_only"})
        self.record(self.monitors[0], "unknown")
        self.record(self.monitors[0], ok=True)
        with self.db._conn:
            self.db._conn.execute("UPDATE email_deliveries SET rule_type_at_creation = ''")
        self.db.close()
        root = Path(self.temp.name)
        self.db = Database(root / "db", SecurityManager(root / "key"))
        self.db.save_email_rule({**self.db.get_email_rule(rule_id), "rule_type": "incident", "excluded_error_types": ["unknown"]})
        notifier = EmailNotifier(self.db)
        with patch.object(notifier, "_send") as send:
            notifier.process_pending()
        send.assert_not_called()

    def test_interrupted_incident_link_backfill_is_resumed(self):
        self.record(self.monitors[0], "http_error")
        with self.db._conn:
            self.db._conn.execute("UPDATE email_incident_state SET down_event_id = NULL")
        self.db.close()
        root = Path(self.temp.name)
        self.db = Database(root / "db", SecurityManager(root / "key"))
        self.db.save_email_rule({**self.db.get_email_rule(self.rule["id"]), "excluded_error_types": ["http_error"]})
        self.record(self.monitors[0], ok=True)
        self.assertEqual(self.db.email_delivery_counts(), {"suppressed": 1})


if __name__ == "__main__":
    unittest.main()
