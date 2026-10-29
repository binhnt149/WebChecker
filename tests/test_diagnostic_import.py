import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sitesentry.core.database import Database
from sitesentry.core.excel_importer import create_monitor_template, import_monitors_from_excel
from sitesentry.core.security import SecurityManager


class DiagnosticImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db = Database(self.root / "db", SecurityManager(self.root / "key"))
        self.db.save_network_profile({"name": "VPN", "scope": "intranet_vpn", "targets": []})
        self.rules = [{"id": "failure", "enabled": True, "priority": 1, "phase": "response", "conditions": [{"source": "json", "path": "success", "operator": "equals", "value": False}], "result": {"error_type": "assertion_failed", "reason_code": "business_rule_match"}}]
        self.id = self.db.save_monitor({"name": "Portal", "url": "https://example.test", "network_profile": "VPN", "diagnostic_rules_json": json.dumps(self.rules)})

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def workbook(self, columns, values):
        path = self.root / "import.xlsx"
        book = Workbook()
        book.active.title = "Monitors"
        book.active.append(columns)
        book.active.append(values)
        book.save(path)
        book.close()
        return path

    def test_old_file_preserves_existing_profile_and_rules(self):
        path = self.workbook(["id", "name", "monitor_type", "url"], [self.id, "Portal", "simple_web", "https://example.test"])
        result = import_monitors_from_excel(self.db, path)
        self.assertEqual(result.updated, 1)
        monitor = self.db.get_monitor(self.id)
        self.assertEqual(monitor["network_profile"], "VPN")
        self.assertEqual(json.loads(monitor["diagnostic_rules_json"]), self.rules)

    def test_blank_preserves_and_explicit_clear_removes(self):
        columns = ["id", "name", "monitor_type", "url", "network_profile", "diagnostic_rules_json"]
        prefix = [self.id, "Portal", "simple_web", "https://example.test"]
        self.assertTrue(import_monitors_from_excel(self.db, self.workbook(columns, prefix + [None, None])).ok)
        self.assertEqual(self.db.get_monitor(self.id)["network_profile"], "VPN")
        self.assertTrue(import_monitors_from_excel(self.db, self.workbook(columns, prefix + ["__CLEAR__", "[]"])).ok)
        monitor = self.db.get_monitor(self.id)
        self.assertEqual(monitor["network_profile"], "")
        self.assertEqual(json.loads(monitor["diagnostic_rules_json"]), [])

    def test_invalid_rows_are_reported_without_partial_update(self):
        columns = ["id", "name", "monitor_type", "url", "network_profile", "diagnostic_rules_json"]
        for tail in [["VPN", "[invalid"], ["Missing", json.dumps(self.rules)]]:
            with self.subTest(tail=tail):
                result = import_monitors_from_excel(self.db, self.workbook(columns, [self.id, "Changed", "api", "https://example.test"] + tail))
                self.assertFalse(result.ok)
                self.assertIn("Row 2", result.errors[0])
                self.assertEqual(self.db.get_monitor(self.id)["name"], "Portal")

    def test_template_exposes_new_columns_and_instructions(self):
        path = self.root / "template.xlsx"
        create_monitor_template(path)
        book = load_workbook(path)
        try:
            headers = [cell.value for cell in book["Monitors"][1]]
            self.assertEqual(headers[-2:], ["network_profile", "diagnostic_rules_json"])
            notes = [row[0].value for row in book["Instructions"].iter_rows()]
            self.assertIn("network_profile", notes)
            self.assertIn("diagnostic_rules_json", notes)
        finally:
            book.close()
